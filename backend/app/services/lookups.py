import unicodedata
from collections.abc import Iterator

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import LOOKUPS_XLSX_PATH
from app.models import LookupHolder, LookupResponse, LookupTankType

HOLDER_SHEET_KEYWORD = "drzitel"
TANK_TYPE_SHEET_KEYWORD = "cistern"

HOLDER_COLUMNS = {
    "holder_name": ("držiteľ", "drzitel"),
    "holder_street": ("ulica",),
    "holder_postal_code": ("psč", "psc"),
    "holder_city": ("mesto",),
    "holder_country": ("štát", "stat"),
}

TANK_TYPE_COLUMNS = {
    "type_approval_number": ("číslo schválenia",),
    "tank_manufacturer_name": ("názov výrobcu",),
    "tank_code": ("kód cisterny",),
    "shell_thickness": ("hrúbka steny cisterny",),
    "head_thickness": ("hrúbka steny dien",),
    "test_pressure": ("skúšobný tlak",),
    "working_pressure": ("pracovný tlak",),
    "calculation_pressure": ("výpočtový pretlak",),
    "calculation_vacuum": ("výpočtový podtlak",),
}

HEADER_SCAN_ROWS = 10


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value: object) -> str:
    return cell_text(value).casefold()


def ascii_fold(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def find_sheet(sheet_names: list[str], keyword: str) -> str | None:
    for name in sheet_names:
        if keyword in ascii_fold(name.casefold()):
            return name
    return None


def find_columns(
    sheet: Worksheet,
    columns: dict[str, tuple[str, ...]],
) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True),
        start=1,
    ):
        headers = [normalized(value) for value in row]
        column_map = {}
        for field, keywords in columns.items():
            for column_index, header in enumerate(headers):
                if header and any(keyword in header for keyword in keywords):
                    column_map[field] = column_index
                    break
        if len(column_map) == len(columns):
            return row_index, column_map
    return None


def parse_rows(
    sheet: Worksheet,
    columns: dict[str, tuple[str, ...]],
    key_field: str,
) -> Iterator[dict[str, str | None]]:
    located = find_columns(sheet, columns)
    if located is None:
        return
    header_row, column_map = located

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = {
            field: cell_text(row[index]) if index < len(row) else ""
            for field, index in column_map.items()
        }
        if not values[key_field]:
            continue
        yield {field: value or None for field, value in values.items()}


def load_lookups() -> LookupResponse:
    if not LOOKUPS_XLSX_PATH.exists():
        return LookupResponse(holders=[], tank_types=[])

    workbook = load_workbook(LOOKUPS_XLSX_PATH, read_only=True, data_only=True)
    try:
        holders: list[LookupHolder] = []
        holder_sheet = find_sheet(workbook.sheetnames, HOLDER_SHEET_KEYWORD)
        if holder_sheet:
            holders = [
                LookupHolder.model_validate(row)
                for row in parse_rows(workbook[holder_sheet], HOLDER_COLUMNS, "holder_name")
            ]

        tank_types: list[LookupTankType] = []
        tank_type_sheet = find_sheet(workbook.sheetnames, TANK_TYPE_SHEET_KEYWORD)
        if tank_type_sheet:
            tank_types = [
                LookupTankType.model_validate(row)
                for row in parse_rows(
                    workbook[tank_type_sheet],
                    TANK_TYPE_COLUMNS,
                    "type_approval_number",
                )
            ]

        return LookupResponse(holders=holders, tank_types=tank_types)
    finally:
        workbook.close()
